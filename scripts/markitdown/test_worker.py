"""Synthetic documents exercise the real pinned SDK without any user data."""

from __future__ import annotations

import importlib.util
import io
import json
from pathlib import Path
import struct
import subprocess
import sys
import tempfile
import unittest
import zipfile

WORKER = Path(__file__).with_name("worker.py")
spec = importlib.util.spec_from_file_location("babel_markitdown_worker", WORKER)
worker = importlib.util.module_from_spec(spec)
spec.loader.exec_module(worker)


def archive_bytes(entries: dict[str, str]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    return buffer.getvalue()


def pdf_bytes(text: str | None) -> bytes:
    stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(data)
    data.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return bytes(data)


def xls_bytes() -> bytes:
    # A valid BIFF2 worksheet needs no third-party writer or fixture download.
    def record(code: int, data: bytes) -> bytes:
        return struct.pack("<HH", code, len(data)) + data

    data = record(0x0009, struct.pack("<HH", 0, 0x0010))
    for row, value in enumerate((b"Babel XLS", b"Local conversion")):
        data += record(0x0004, struct.pack("<HH3sB", row, 0, b"\x00\x00\x00", len(value)) + value)
    return data + record(0x000A, b"")


def epub_bytes(chapters: dict[str, str], toc: str | None = None, ncx: bool = False, href: str | None = None) -> bytes:
    items = []
    spine = []
    entries = {"mimetype": "application/epub+zip", "META-INF/container.xml":
               '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OPS/content.opf"/></rootfiles></container>'}
    for index, (path, html) in enumerate(chapters.items()):
        items.append(f'<item id="c{index}" href="{href or path}" media-type="application/xhtml+xml"/>')
        spine.append(f'<itemref idref="c{index}"/>')
        entries[f"OPS/{path}"] = html
    if toc is not None:
        toc_path = "toc.ncx" if ncx else "nav.xhtml"
        media_type = "application/x-dtbncx+xml" if ncx else "application/xhtml+xml"
        items.append(f'<item id="toc" href="{toc_path}" media-type="{media_type}" properties="nav"/>' if not ncx else
                     f'<item id="toc" href="{toc_path}" media-type="{media_type}"/>')
        entries[f"OPS/{toc_path}"] = toc
    entries["OPS/content.opf"] = '<package xmlns="http://www.idpf.org/2007/opf"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Test book</dc:title><dc:creator>Test author</dc:creator></metadata><manifest>' + "".join(items) + '</manifest><spine toc="toc">' + "".join(spine) + '</spine></package>'
    return archive_bytes(entries)


class WorkerTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.directory = tempfile.TemporaryDirectory(prefix="babel-markitdown-tests-")
        cls.root = Path(cls.directory.name)

    @classmethod
    def tearDownClass(cls):
        cls.directory.cleanup()

    def invoke(self, name: str, data: bytes, expected_error: str | None = None) -> dict:
        path = self.root / name
        path.write_bytes(data)
        completed = subprocess.run(
            [sys.executable, "-I", str(WORKER), str(path), path.suffix],
            capture_output=True, encoding="utf-8", timeout=60,
        )
        result = json.loads(completed.stdout)
        self.assertEqual(completed.returncode, 1 if expected_error else 0, result)
        if expected_error:
            self.assertEqual(result["error"]["code"], expected_error)
            self.assertNotIn(str(path), result["error"]["message"])
        else:
            self.assertEqual(result["sdkVersion"], worker.SDK_VERSION)
            self.assertIsInstance(result["warnings"], list)
            self.assertTrue(result["markdown"].strip())
        return result

    def test_probe(self):
        completed = subprocess.run([sys.executable, "-I", str(WORKER), "--probe"], capture_output=True, encoding="utf-8", timeout=30)
        self.assertEqual(completed.returncode, 0, completed.stderr)
        self.assertEqual(json.loads(completed.stdout), {"ready": True, "sdkVersion": worker.SDK_VERSION})

    def test_utf8_text_and_structured_formats(self):
        documents = {
            "sample.txt": "Babel 本地转换\nSecond paragraph",
            "sample.csv": "Name,Value\nBabel,42\n",
            "sample.json": '{"title":"Babel", "value":42}',
            "sample.xml": '<document><title>Babel</title><value>42</value></document>',
        }
        for name, text in documents.items():
            with self.subTest(name=name):
                result = self.invoke(name, text.encode("utf-8"))
                self.assertIn("Babel", result["markdown"])
                if name.endswith((".json", ".xml")):
                    language = Path(name).suffix[1:]
                    self.assertEqual(result["markdown"], f"```{language}\n{text}\n```")
        literal_xml = '<document><img src="file:///private.png"/><code>```</code></document>'
        structured = self.invoke("literal.xml", literal_xml.encode())["markdown"]
        self.assertEqual(structured, f"````xml\n{literal_xml}\n````")
        self.assertIn("本地转换", self.invoke("unicode.txt", "Babel 本地转换".encode())["markdown"])

    def test_html_extracts_text_without_images_or_scripts(self):
        html = b'<html><head><title>Babel import</title></head><body><h1>Babel</h1><script>SECRET_SCRIPT</script><p>Local import</p><img src="https://invalid.example/track"><img src="file:///C:/secret.txt"></body></html>'
        result = self.invoke("sample.html", html)
        self.assertEqual(result["title"], "Babel import")
        self.assertIn("# Babel", result["markdown"])
        self.assertNotIn("SECRET_SCRIPT", result["markdown"])
        self.assertNotIn("invalid.example", result["markdown"])
        self.assertNotIn("secret.txt", result["markdown"])

    def test_docx(self):
        data = archive_bytes({
            "[Content_Types].xml": '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
            "_rels/.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>',
            "word/document.xml": '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Babel DOCX 本地转换</w:t></w:r></w:p></w:body></w:document>',
        })
        self.assertIn("Babel DOCX 本地转换", self.invoke("sample.docx", data)["markdown"])

    def test_pptx(self):
        from pptx import Presentation
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Babel PPTX"
        slide.placeholders[1].text = "Local document conversion"
        buffer = io.BytesIO()
        presentation.save(buffer)
        self.assertIn("Babel PPTX", self.invoke("sample.pptx", buffer.getvalue())["markdown"])

    def test_xlsx(self):
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.append(["Babel XLSX", "Value"])
        workbook.active.append(["Local conversion", 42])
        buffer = io.BytesIO()
        workbook.save(buffer)
        result = self.invoke("sample.xlsx", buffer.getvalue())
        self.assertIn("Babel XLSX", result["markdown"])
        self.assertIn("42", result["markdown"])

    def test_xls(self):
        self.assertIn("Babel XLS", self.invoke("sample.xls", xls_bytes())["markdown"])

    def test_pdf_and_scanned_pdf(self):
        result = self.invoke("sample.pdf", pdf_bytes("Babel PDF local conversion"))
        self.assertIn("Babel PDF", result["markdown"])
        self.assertTrue(any("OCR" in message for message in result["warnings"]))
        self.invoke("scan.pdf", pdf_bytes(None), "no_extractable_text")

    def test_epub(self):
        data = archive_bytes({
            "mimetype": "application/epub+zip",
            "META-INF/container.xml": '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>',
            "OPS/content.opf": '<package xmlns="http://www.idpf.org/2007/opf"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Babel EPUB</dc:title></metadata><manifest><item id="chapter" href="chapter.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="chapter"/></spine></package>',
            "OPS/chapter.xhtml": '<html xmlns="http://www.w3.org/1999/xhtml"><body><h1>Babel chapter</h1><p>Local EPUB conversion.</p></body></html>',
        })
        result = self.invoke("sample.epub", data)
        self.assertEqual(result["title"], "Babel EPUB")
        self.assertIn("Babel chapter", result["markdown"])
        self.assertEqual(result["chapterSource"], "epub-spine")
        self.assertEqual(result["chapters"][0]["title"], "Babel chapter")

    def test_epub_toc_uses_spine_order_and_omits_parent_path_images_before_markdown(self):
        toc = '<html><body><nav epub:type="toc"><ol><li><a href="Text/last.xhtml">Second chapter</a></li><li><a href="Text/first.xhtml#start">First chapter</a></li></ol></nav></body></html>'
        data = epub_bytes({
            "Text/first.xhtml": '<html><body><h1 id="start">Start</h1><p>First section text.</p><img src="../images/p031_01.png" alt="Plot [31] nested ]" title="Line\nTitle"/><img src="https://invalid.example/tracker.png"/><svg><image href="../images/svg.png"/></svg></body></html>',
            "Text/last.xhtml": '<html><body><h1>Finish</h1><p>Second section text.</p></body></html>',
        }, toc)
        result = self.invoke("book-with-images.epub", data)
        self.assertEqual(result["chapterSource"], "epub-toc")
        self.assertEqual([chapter["title"] for chapter in result["chapters"]], ["First chapter", "Second chapter"])
        self.assertIn("**Authors:** Test author", result["markdown"])
        for markdown in [result["markdown"], *(chapter["markdown"] for chapter in result["chapters"])]:
            self.assertNotIn("../images", markdown)
            self.assertNotIn("![", markdown)
            self.assertNotIn("invalid.example", markdown)
        self.assertTrue(any("Images" in warning for warning in result["warnings"]))

    def test_epub_toc_splits_anchors_in_a_single_file_without_losing_introductory_text(self):
        toc = '<html><body><nav role="doc-toc"><ol><li><a href="book.xhtml#one">第一章</a></li><li><a href="book.xhtml#one">Duplicate</a></li><li><a href="book.xhtml#two">第二章</a></li></ol></nav></body></html>'
        html = '<html><head><title>Introduction</title></head><body><p>Introductory text.</p><h1 id="one">One</h1><p>First content.</p><h1 id="two">Two</h1><p>Second content.</p></body></html>'
        result = self.invoke("anchored-book.epub", epub_bytes({"book.xhtml": html}, toc))
        self.assertEqual([chapter["title"] for chapter in result["chapters"]], ["Test book — Introduction", "第一章", "第二章"])
        self.assertIn("Introductory text", result["chapters"][0]["markdown"])
        self.assertIn("# One", result["chapters"][1]["markdown"])
        self.assertIn("First content", result["chapters"][1]["markdown"])
        self.assertNotIn("Second content", result["chapters"][1]["markdown"])
        self.assertIn("Second content", result["chapters"][2]["markdown"])
        self.assertNotIn("BABELCHAPTER", result["markdown"])

    def test_epub_ncx_titles_and_relative_archive_paths(self):
        toc = '<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/"><navMap><navPoint><navLabel><text>NCX chapter</text></navLabel><content src="Text/one.xhtml"/></navPoint></navMap></ncx>'
        result = self.invoke("ncx-book.epub", epub_bytes({"Text/one.xhtml": "<h1>Heading</h1><p>NCX content.</p>"}, toc, ncx=True, href="Text/../Text/one.xhtml"))
        self.assertEqual(result["chapterSource"], "epub-toc")
        self.assertEqual(result["chapters"][0]["title"], "NCX chapter")
        self.assertIn("NCX content", result["chapters"][0]["markdown"])

    def test_epub_content_references_cannot_escape_the_archive_or_request_external_files(self):
        for href in ("../../outside.xhtml", "%2e%2e/%2e%2e/outside.xhtml", "file:///private.xhtml", "https://invalid.example/book.xhtml", "/absolute.xhtml"):
            with self.subTest(href=href):
                self.invoke("unsafe-content.epub", epub_bytes({"one.xhtml": "<h1>Text</h1>"}, href=href), "unsafe_archive")

    def test_html_nested_image_alts_and_literal_code_examples(self):
        result = self.invoke("nested-alt.html", b'<h1>Text</h1><img src="../images/p031_01.png" alt="Figure [31] ]"/><pre>![code](../images/example.png)</pre>')
        self.assertNotIn("p031_01", result["markdown"])
        self.assertIn("![code](../images/example.png)", result["markdown"])

    def test_invalid_empty_and_unsupported_documents(self):
        self.invoke("empty.txt", b"", "empty_document")
        self.invoke("sample.exe", b"not allowed", "unsupported_format")
        self.invoke("broken.docx", b"broken archive", "invalid_document")
        self.invoke("protected.docx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"x" * 20, "encrypted_document")
        self.invoke("broken.pdf", b"not a PDF", "conversion_failed")

    def test_zip_bombs_and_unsafe_paths(self):
        for filename in ("../outside.xml", "/absolute.xml", "C:/outside.xml"):
            with self.subTest(filename=filename):
                self.invoke("unsafe.docx", archive_bytes({filename: "Babel"}), "unsafe_archive")
        self.invoke("bomb.docx", archive_bytes({"word/document.xml": "0" * (2 * 1024 * 1024)}), "unsafe_archive")
        self.invoke("disguised.txt", archive_bytes({"text.txt": "Babel"}), "invalid_document")

    def test_encrypted_zip_is_rejected(self):
        data = bytearray(archive_bytes({"word/document.xml": "Babel"}))
        for signature, flag_offset in ((b"PK\x03\x04", 6), (b"PK\x01\x02", 8)):
            start = data.index(signature)
            data[start + flag_offset] |= 1
        self.invoke("encrypted.docx", bytes(data), "encrypted_document")

    def test_input_and_output_size_limits(self):
        self.invoke("large.txt", b"x" * (worker.MAX_INPUT_BYTES + 1), "input_too_large")
        self.invoke("large-output.txt", (b"Babel text line\n" * 700000), "output_too_large")

    def test_external_io_policy(self):
        for event in ("socket.connect", "socket.getaddrinfo", "subprocess.Popen", "os.system"):
            with self.subTest(event=event), self.assertRaises(worker.WorkerError):
                worker.block_external_io(event, ())
        with self.assertRaises(worker.WorkerError):
            worker.LocalOnlySession().get("https://invalid.example")

    def test_sdk_network_requests_are_blocked_in_the_worker_process(self):
        path = self.root / "network.html"
        path.write_text("<p>Babel network check</p>", encoding="utf-8")
        code = """
import json, runpy, sys
module = runpy.run_path(sys.argv[1])
from markitdown import FileConversionException
from markitdown.converters import HtmlConverter
import requests
def attempt_network(self, *args, **kwargs):
    requests.get('http://127.0.0.1:1/', timeout=1)
HtmlConverter.convert = attempt_network
try:
    module['convert'](sys.argv[2], '.html')
    raise AssertionError('Conversion unexpectedly permitted a network request')
except FileConversionException as error:
    exceptions = [attempt.exc_info[1] for attempt in error.attempts]
    assert exceptions and all(getattr(item, 'code', None) == 'external_access_blocked' for item in exceptions), exceptions
    print(json.dumps({'blocked': True}))
"""
        completed = subprocess.run([sys.executable, "-I", "-c", code, str(WORKER), str(path)], capture_output=True, encoding="utf-8", timeout=30)
        self.assertEqual(completed.returncode, 0, completed.stderr)
        self.assertEqual(json.loads(completed.stdout), {"blocked": True})


if __name__ == "__main__":
    unittest.main(verbosity=2)
